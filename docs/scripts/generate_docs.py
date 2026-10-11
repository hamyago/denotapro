#!/usr/bin/env python3
"""
Générateur de documents DeNoTa
Convertit tous les fichiers Markdown en PDF et DOCX.

Prérequis :
    pip install python-docx markdown beautifulsoup4 weasyprint md2pdf

Usage :
    python generate_docs.py
"""

import sys
from pathlib import Path
from markdown import markdown
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ── Configuration ─────────────────────────────────────────
SCRIPT_DIR = Path(__file__).parent
DOCS_DIR = SCRIPT_DIR.parent
OUTPUT_DIR = DOCS_DIR / "output"

# Couleurs DeNoTa
GREEN = RGBColor(0x1B, 0x5E, 0x3B)
GOLD = RGBColor(0xF5, 0xA6, 0x23)
INK = RGBColor(0x0D, 0x1B, 0x2A)
GREY = RGBColor(0x6B, 0x77, 0x85)

# Fichiers à convertir
FILES = [
    ("CONTEXT_DENOTA_v1.2.md", "Contexte_DeNoTa_v1.2"),
    ("institutionnel/01_DOSSIER_MINISTERE.md", "01_Dossier_Ministere"),
    ("institutionnel/02_PITCH_5MIN.md", "02_Pitch_5min"),
    ("institutionnel/03_PROTOCOLE_ACCORD.md", "03_Protocole_Accord"),
    ("institutionnel/04_COMMUNIQUE_PRESSE.md", "04_Communique_Presse"),
    ("institutionnel/05_SCRIPT_AUDIENCE.md", "05_Script_Audience"),
    ("institutionnel/06_NOTE_STRATEGIQUE_SIGS_JTS.md", "06_Note_Strategique_SIGS_JTS"),
    ("institutionnel/07_ARGUMENTAIRE_FONTSI_ANSUT.md", "07_Argumentaire_FONTSI_ANSUT"),
]


# ── Utilitaires ───────────────────────────────────────────
def ensure_output_dir():
    """Crée le dossier de sortie."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def clean_markdown(text: str) -> str:
    """Nettoie le markdown inline pour l'affichage DOCX."""
    text = text.replace('**', '').replace('__', '')
    text = text.replace('*', '').replace('_', '')
    text = text.replace('`', '')
    text = text.replace('> ', '')
    return text


# ── Conversion DOCX ───────────────────────────────────────
def md_to_docx(md_path: Path, docx_path: Path):
    """Convertit un fichier Markdown en DOCX avec styles DeNoTa."""
    doc = Document()

    # Marges
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    # Style de base
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')

    for line in lines:
        line = line.rstrip()

        if not line.strip():
            continue

        # Titres
        if line.startswith('# '):
            h = doc.add_heading(line[2:], level=1)
            for run in h.runs:
                run.font.color.rgb = GREEN
        elif line.startswith('## '):
            h = doc.add_heading(line[3:], level=2)
            for run in h.runs:
                run.font.color.rgb = GREEN
        elif line.startswith('### '):
            h = doc.add_heading(line[4:], level=3)
            for run in h.runs:
                run.font.color.rgb = GOLD
        elif line.startswith('#### '):
            h = doc.add_heading(line[5:], level=4)
            for run in h.runs:
                run.font.color.rgb = INK

        # Citations
        elif line.startswith('> '):
            p = doc.add_paragraph(line[2:])
            p.paragraph_format.left_indent = Cm(1)
            for run in p.runs:
                run.font.italic = True
                run.font.color.rgb = GREEN

        # Listes
        elif line.startswith('- ') or line.startswith('* '):
            doc.add_paragraph(clean_markdown(line[2:]), style='List Bullet')
        elif line.startswith('  - ') or line.startswith('  * '):
            doc.add_paragraph(clean_markdown(line[4:]), style='List Bullet 2')

        # Séparateurs
        elif line.startswith('---'):
            doc.add_paragraph('─' * 60)

        # Tableaux
        elif line.startswith('|'):
            if '---' in line:
                continue
            cells = [c.strip() for c in line.split('|')[1:-1]]
            if cells:
                p = doc.add_paragraph()
                p.add_run(' | '.join(cells))

        # Paragraphes normaux
        else:
            text = clean_markdown(line)
            if text.strip():
                doc.add_paragraph(text)

    doc.save(docx_path)
    print(f"✅ DOCX généré : {docx_path.name}")


# ── Conversion PDF via WeasyPrint ────────────────────────
def md_to_pdf(md_path: Path, pdf_path: Path):
    """Convertit un fichier Markdown en PDF via WeasyPrint."""
    try:
        from weasyprint import HTML, CSS

        with open(md_path, 'r', encoding='utf-8') as f:
            md_content = f.read()

        # Convertir Markdown → HTML
        html_body = markdown(
            md_content,
            extensions=['extra', 'tables', 'toc', 'sane_lists']
        )

        # CSS DeNoTa
        css_style = """
        @page {
            size: A4;
            margin: 2cm;
            @bottom-center {
                content: "DeNoTa — " counter(page);
                font-family: 'Helvetica', sans-serif;
                font-size: 9pt;
                color: #6B7785;
            }
        }
        body {
            font-family: 'Helvetica', 'Arial', sans-serif;
            font-size: 11pt;
            line-height: 1.6;
            color: #0D1B2A;
        }
        h1 {
            color: #1B5E3B;
            font-size: 22pt;
            border-bottom: 2px solid #F5A623;
            padding-bottom: 6px;
            margin-top: 20px;
        }
        h2 {
            color: #1B5E3B;
            font-size: 16pt;
            margin-top: 18px;
        }
        h3 {
            color: #F5A623;
            font-size: 13pt;
            margin-top: 14px;
        }
        h4 {
            color: #0D1B2A;
            font-size: 12pt;
        }
        table {
            border-collapse: collapse;
            width: 100%;
            margin: 10px 0;
        }
        th, td {
            border: 1px solid #E8ECF0;
            padding: 6px 10px;
            text-align: left;
            font-size: 10pt;
        }
        th {
            background-color: #1B5E3B;
            color: white;
        }
        tr:nth-child(even) {
            background-color: #F8FAFC;
        }
        blockquote {
            border-left: 4px solid #F5A623;
            padding-left: 12px;
            color: #1B5E3B;
            font-style: italic;
            margin: 10px 0;
        }
        code {
            background-color: #F1F5F9;
            padding: 2px 4px;
            border-radius: 3px;
            font-size: 9pt;
        }
        ul, ol {
            margin: 8px 0;
        }
        hr {
            border: none;
            border-top: 1px solid #E8ECF0;
            margin: 16px 0;
        }
        """

        html_content = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="UTF-8">
            <style>{css_style}</style>
        </head>
        <body>
            {html_body}
        </body>
        </html>
        """

        HTML(string=html_content).write_pdf(str(pdf_path))
        print(f"✅ PDF généré  : {pdf_path.name}")

    except ImportError:
        print(f"⚠️  WeasyPrint non installé. Installez-le : pip install weasyprint")
    except Exception as e:
        print(f"❌ Erreur PDF {md_path.name} : {e}")


# ── Main ──────────────────────────────────────────────────
def main():
    print("🚀 Génération des documents DeNoTa...\n")
    ensure_output_dir()

    for md_file, base_name in FILES:
        md_path = DOCS_DIR / md_file
        if not md_path.exists():
            print(f"⚠️  Fichier introuvable : {md_path}")
            continue

        print(f"📄 Traitement : {md_file}")

        # DOCX
        docx_path = OUTPUT_DIR / f"{base_name}.docx"
        try:
            md_to_docx(md_path, docx_path)
        except Exception as e:
            print(f"❌ Erreur DOCX {md_file} : {e}")

        # PDF
        pdf_path = OUTPUT_DIR / f"{base_name}.pdf"
        md_to_pdf(md_path, pdf_path)

        print()

    print(f"✨ Terminé ! Documents dans : {OUTPUT_DIR}")


if __name__ == "__main__":
    main()